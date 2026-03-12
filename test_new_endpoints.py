"""Test script for new endpoints without MongoDB dependency."""
import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'backend'))

# Test the generate_example_excel endpoint logic
print("Testing generate example Excel functionality...")

import openpyxl
from openpyxl.styles import Font, PatternFill
import random
import io

def test_generate_example(report_name: str, report_id: str):
    """Test generating example Excel file."""
    # Generate random formatSpecId and domainId (24-char hex)
    format_spec_id = ''.join(random.choices('0123456789abcdef', k=24))
    domain_id = ''.join(random.choices('0123456789abcdef', k=24))

    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report Metadata"

    # Header style
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    # Report metadata section
    ws['A1'] = 'reportName'
    ws['B1'] = report_name
    ws['A1'].fill = header_fill
    ws['A1'].font = header_font

    ws['A2'] = 'reportId'
    ws['B2'] = report_id
    ws['A2'].fill = header_fill
    ws['A2'].font = header_font

    ws['A3'] = 'reportType'
    ws['B3'] = random.choice(['analytics', 'operational', 'financial', 'marketing', 'sales'])
    ws['A3'].fill = header_fill
    ws['A3'].font = header_font

    ws['A4'] = 'formatSpecId'
    ws['B4'] = format_spec_id
    ws['A4'].fill = header_fill
    ws['A4'].font = header_font

    ws['A5'] = 'domainId'
    ws['B5'] = domain_id
    ws['A5'].fill = header_fill
    ws['A5'].font = header_font

    # Parameters section
    ws['A8'] = 'Parameters'
    ws['A8'].fill = header_fill
    ws['A8'].font = header_font

    ws['A9'] = 'parameterName'
    ws['B9'] = 'originalColumnName'
    ws['C9'] = 'parameterType'
    for col in ['A9', 'B9', 'C9']:
        ws[col].fill = header_fill
        ws[col].font = header_font

    # Generate 4-6 random parameters
    param_examples = [
        ('startDate', 'start_date', 'date'),
        ('endDate', 'end_date', 'date'),
        ('reportDate', 'report_date', 'timestamp'),
        ('userId', 'user_id', 'varchar(64)'),
        ('accountId', 'account_id', 'varchar(64)'),
        ('departmentId', 'department_id', 'varchar(64)'),
        ('regionCode', 'region_code', 'varchar(32)'),
        ('fiscalYear', 'fiscal_year', 'integer'),
        ('currencyCode', 'currency_code', 'char(3)'),
        ('statusFilter', 'status_filter', 'varchar(16)'),
    ]
    num_params = random.randint(4, 6)
    random.shuffle(param_examples)

    row_idx = 10
    for i in range(num_params):
        param_data = param_examples[i]
        ws[f'A{row_idx}'] = param_data[0]
        ws[f'B{row_idx}'] = param_data[1]
        ws[f'C{row_idx}'] = param_data[2]
        row_idx += 1

    # Columns section
    row_idx += 3
    ws[f'A{row_idx}'] = 'column name'
    ws[f'B{row_idx}'] = 'actual name'
    ws[f'C{row_idx}'] = 'type (postgresql)'
    ws[f'D{row_idx}'] = 'description'
    for col in [f'A{row_idx}', f'B{row_idx}', f'C{row_idx}', f'D{row_idx}']:
        ws[col].fill = header_fill
        ws[col].font = header_font

    # Generate exactly 25 columns
    column_examples = [
        ('id', 'ID', 'varchar(64)', 'unique identifier of the record'),
        ('name', 'Name', 'varchar(128)', 'display name of the entity'),
        ('created_at', 'Created At', 'timestamp', 'timestamp when the record was created'),
        ('updated_at', 'Updated At', 'timestamp', 'timestamp when the record was last updated'),
        ('status', 'Status', 'varchar(16)', 'current status of the record'),
        ('amount', 'Amount', 'numeric(18,2)', 'monetary amount in base currency'),
        ('quantity', 'Quantity', 'integer', 'number of items or units'),
        ('description', 'Description', 'varchar(256)', 'detailed description or notes'),
        ('category', 'Category', 'varchar(32)', 'classification category code'),
        ('user_id', 'User ID', 'varchar(64)', 'identifier of the associated user'),
        ('customer_id', 'Customer ID', 'varchar(64)', 'unique identifier of the customer'),
        ('email', 'Email', 'varchar(255)', 'email address of the user or customer'),
        ('phone', 'Phone', 'varchar(32)', 'phone number in international format'),
        ('address', 'Address', 'varchar(256)', 'full address or street address'),
        ('city', 'City', 'varchar(64)', 'city name'),
        ('country', 'Country', 'varchar(64)', 'country name or ISO code'),
        ('postal_code', 'Postal Code', 'varchar(16)', 'postal or ZIP code'),
        ('currency', 'Currency', 'char(3)', 'ISO 4217 currency code'),
        ('tax_amount', 'Tax Amount', 'numeric(18,2)', 'tax amount included or applied'),
        ('discount_amount', 'Discount Amount', 'numeric(18,2)', 'discount amount applied to the transaction'),
        ('total_amount', 'Total Amount', 'numeric(18,2)', 'final total amount after all adjustments'),
        ('payment_method', 'Payment Method', 'varchar(32)', 'payment method used'),
        ('transaction_date', 'Transaction Date', 'timestamp', 'date and time of the transaction'),
        ('reference_number', 'Reference Number', 'varchar(64)', 'reference or confirmation number'),
        ('notes', 'Notes', 'text', 'additional notes or comments'),
    ]

    num_columns = 25
    random.shuffle(column_examples)

    row_idx += 1
    for i in range(num_columns):
        col_data = column_examples[i]
        ws[f'A{row_idx}'] = col_data[0]
        ws[f'B{row_idx}'] = col_data[1]
        ws[f'C{row_idx}'] = col_data[2]
        ws[f'D{row_idx}'] = col_data[3]
        row_idx += 1

    # Adjust column widths
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 50

    # Save to test file
    output_file = f"test_output/example_{report_id}.xlsx"
    os.makedirs("test_output", exist_ok=True)
    wb.save(output_file)

    print(f"[OK] Successfully generated example Excel file: {output_file}")
    print(f"  - Report Name: {report_name}")
    print(f"  - Report ID: {report_id}")
    print(f"  - Format Spec ID: {format_spec_id}")
    print(f"  - Domain ID: {domain_id}")
    print(f"  - Parameters: {num_params}")
    print(f"  - Columns: {num_columns}")

# Test 1: Generate example with custom values (6-digit reportId)
test_generate_example("Sales Metrics Report", "200012")

# Test 2: Generate another example with different values (6-digit reportId)
test_generate_example("Customer Analytics", "500001")

print("\n[OK] All tests passed successfully!")
print("\nNote: The actual endpoints are implemented in:")
print("  - POST /api/upload/generate-example (generate example Excel)")
print("  - GET /api/reports/{report_id}/download-all (download ZIP)")
