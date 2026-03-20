"""Upload and processing endpoints."""

from fastapi import APIRouter, UploadFile, File, HTTPException, BackgroundTasks, Request
from fastapi.responses import JSONResponse, StreamingResponse
from typing import Dict, Tuple
from pydantic import BaseModel
import tempfile
import os
import time
import logging
from pathlib import Path
import uuid
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import random
import io

from app.limiter import limiter
from app.services import ExcelParser, LangChainProcessor, db_service, ReportValidator
from app.models.migration_log import MigrationLog
from app.models.report import Report
from app.models.role import Role
from app.models.report_config import ReportConfig
from app.models.report_ui_settings import ReportUiSettings
from app.config import settings
import logging

router = APIRouter(prefix="/api/upload", tags=["upload"])
logger = logging.getLogger(__name__)

# Initialize processors
langchain_processor = LangChainProcessor()
report_validator = ReportValidator()

# Temporary storage for pending uploads awaiting user confirmation
# Key: session_id, Value: (report, roles, config, ui_settings, file_info)
pending_uploads: Dict[str, Tuple[Report, list, ReportConfig, ReportUiSettings, Dict]] = {}


class ConfirmUploadRequest(BaseModel):
    """Request to confirm and save a pending upload."""
    sessionId: str


class GenerateExampleRequest(BaseModel):
    """Request to generate an example Excel file."""
    reportName: str
    reportId: str


@router.post("")
@limiter.limit("5/minute")
async def upload_excel(
    request: Request,
    file: UploadFile = File(...),
    background_tasks: BackgroundTasks = None
) -> Dict:
    """
    Upload and process Excel file containing report metadata.

    Process:
    1. Validate file type and size
    2. Parse Excel file → report, roles, config, ui_settings
    3. Enhance with AI (LangChain + Claude)
    4. Save everything to MongoDB
    5. Record migration log

    Returns:
        Processing result with report details
    """
    # Validate file type
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(
            status_code=400,
            detail="Invalid file type. Only Excel files (.xlsx, .xls) are allowed."
        )

    # Read file content
    content = await file.read()
    file_size_mb = len(content) / (1024 * 1024)

    # Validate file size
    if file_size_mb > settings.max_file_size_mb:
        raise HTTPException(
            status_code=400,
            detail=f"File too large. Maximum size is {settings.max_file_size_mb}MB."
        )

    logger.info(f"Processing upload: {file.filename} ({file_size_mb:.2f}MB)")

    start_time = time.time()
    temp_file = None

    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp:
            tmp.write(content)
            temp_file = tmp.name

        # ── Step 1: Parse Excel file ──
        parser = ExcelParser(temp_file)
        report, roles, report_config, ui_settings = parser.parse()
        parser.close()

        logger.info(f"Parsed report {report.reportId}: {len(report.columns)} columns")

        # ── Step 2: Enhance with LangChain ──
        ai_enhanced = False
        try:
            report = await langchain_processor.enhance_report(report)
            ai_enhanced = True
            logger.info(f"Enhanced report {report.reportId} with AI")
        except Exception as e:
            logger.error(f"AI enhancement failed: {e}. Continuing with basic data.")

        # ── Step 3: Validate report (deterministic + AI semantic) ──
        validation_result = await report_validator.validate(report)
        validation = validation_result.model_dump()

        # ── Decision Point: Check if there are errors or warnings ──
        has_issues = validation_result.errorCount > 0 or validation_result.warningCount > 0

        if has_issues:
            # Store pending upload and return for user confirmation
            session_id = str(uuid.uuid4())
            file_info = {
                "filename": file.filename,
                "fileSizeMb": file_size_mb,
                "processingTimeMs": (time.time() - start_time) * 1000,
                "aiEnhanced": ai_enhanced,
            }
            pending_uploads[session_id] = (report, roles, report_config, ui_settings, file_info)

            logger.info(
                f"Report {report.reportId} has {validation_result.errorCount} errors "
                f"and {validation_result.warningCount} warnings — awaiting user confirmation "
                f"(session: {session_id})"
            )

            return {
                "success": False,
                "pending": True,
                "sessionId": session_id,
                "reportId": report.reportId,
                "reportName": report.reportName,
                "reportType": report.reportType,
                "columnsCount": len(report.columns),
                "parametersCount": len(report.parameters),
                "rolesCount": len(roles),
                "validation": validation,
                "message": f"Validation found {validation_result.errorCount} error(s) and "
                           f"{validation_result.warningCount} warning(s). Please review and confirm."
            }

        # ── Step 4: No issues — Save everything to MongoDB directly ──
        report_id = await db_service.save_report(report)
        roles_count = await db_service.save_roles(report_id, roles)
        await db_service.save_report_config(report_config)
        await db_service.save_report_ui_settings(ui_settings)

        processing_time_ms = (time.time() - start_time) * 1000

        # ── Step 5: Record migration log (success) ──
        migration_log = MigrationLog(
            reportId=report_id,
            reportName=report.reportName,
            filename=file.filename,
            status="success",
            columnsCount=len(report.columns),
            parametersCount=len(report.parameters),
            rolesCount=roles_count,
            validation=validation,
            fileSizeMb=round(file_size_mb, 4),
            processingTimeMs=round(processing_time_ms, 2),
            aiEnhanced=ai_enhanced,
        )
        await db_service.save_migration_log(migration_log)

        logger.info(
            f"Saved report {report_id} with {roles_count} roles, "
            f"config, UI settings, and migration log "
            f"({processing_time_ms:.0f}ms)"
        )

        return {
            "success": True,
            "pending": False,
            "reportId": report_id,
            "reportName": report.reportName,
            "reportType": report.reportType,
            "columnsCount": len(report.columns),
            "parametersCount": len(report.parameters),
            "rolesCount": roles_count,
            "validation": validation,
            "message": f"Successfully processed {file.filename}"
        }

    except Exception as e:
        processing_time_ms = (time.time() - start_time) * 1000
        logger.error(f"Error processing file: {e}", exc_info=True)

        # Record migration log (failure)
        try:
            fail_log = MigrationLog(
                reportId="unknown",
                reportName="",
                filename=file.filename or "",
                status="failed",
                errorMessage=str(e),
                fileSizeMb=round(file_size_mb, 4),
                processingTimeMs=round(processing_time_ms, 2),
            )
            await db_service.save_migration_log(fail_log)
        except Exception as log_err:
            logger.warning(f"Failed to save error migration log: {log_err}")

        raise HTTPException(
            status_code=500,
            detail=f"Error processing file: {str(e)}"
        )

    finally:
        # Clean up temporary file
        if temp_file and os.path.exists(temp_file):
            try:
                os.unlink(temp_file)
            except Exception as e:
                logger.warning(f"Failed to delete temp file: {e}")


@router.post("/confirm")
async def confirm_upload(request: ConfirmUploadRequest) -> Dict:
    """
    Confirm and save a pending upload after user reviews validation issues.

    Args:
        request: Contains sessionId for the pending upload

    Returns:
        Final save result with report details
    """
    session_id = request.sessionId

    # Retrieve pending upload
    if session_id not in pending_uploads:
        raise HTTPException(
            status_code=404,
            detail=f"No pending upload found for session {session_id}. It may have expired."
        )

    report, roles, report_config, ui_settings, file_info = pending_uploads.pop(session_id)

    start_time = time.time()

    try:
        # Save everything to MongoDB
        report_id = await db_service.save_report(report)
        roles_count = await db_service.save_roles(report_id, roles)
        await db_service.save_report_config(report_config)
        await db_service.save_report_ui_settings(ui_settings)

        processing_time_ms = (time.time() - start_time) * 1000

        # Re-validate to get current state (in case user fixed issues via Quick Fix)
        validation_result = await report_validator.validate(report)
        validation = validation_result.model_dump()

        # Record migration log (success)
        migration_log = MigrationLog(
            reportId=report_id,
            reportName=report.reportName,
            filename=file_info["filename"],
            status="success",
            columnsCount=len(report.columns),
            parametersCount=len(report.parameters),
            rolesCount=roles_count,
            validation=validation,
            fileSizeMb=round(file_info["fileSizeMb"], 4),
            processingTimeMs=round(file_info["processingTimeMs"] + processing_time_ms, 2),
            aiEnhanced=file_info["aiEnhanced"],
        )
        await db_service.save_migration_log(migration_log)

        logger.info(
            f"Confirmed and saved report {report_id} with {roles_count} roles "
            f"(session: {session_id})"
        )

        return {
            "success": True,
            "reportId": report_id,
            "reportName": report.reportName,
            "reportType": report.reportType,
            "columnsCount": len(report.columns),
            "parametersCount": len(report.parameters),
            "rolesCount": roles_count,
            "validation": validation,
            "message": f"Successfully saved {file_info['filename']} after confirmation"
        }

    except Exception as e:
        # Put it back if save fails
        pending_uploads[session_id] = (report, roles, report_config, ui_settings, file_info)
        logger.error(f"Error saving confirmed upload: {e}", exc_info=True)
        raise HTTPException(
            status_code=500,
            detail=f"Error saving report: {str(e)}"
        )


@router.post("/batch")
@limiter.limit("3/minute")
async def upload_batch(
    request: Request,
    files: list[UploadFile] = File(...)
) -> Dict:
    """
    Upload and process multiple Excel files in batch.

    Args:
        files: List of Excel files

    Returns:
        Batch processing results
    """
    if len(files) > 10:
        raise HTTPException(
            status_code=400,
            detail="Maximum 10 files per batch"
        )

    results = []
    errors = []

    for file in files:
        try:
            # Process each file
            result = await upload_excel(file)
            results.append({
                "filename": file.filename,
                "status": "success",
                "reportId": result["reportId"]
            })
        except Exception as e:
            errors.append({
                "filename": file.filename,
                "status": "error",
                "error": str(e)
            })

    return {
        "total": len(files),
        "successful": len(results),
        "failed": len(errors),
        "results": results,
        "errors": errors
    }


@router.post("/generate-example")
async def generate_example_excel(request: GenerateExampleRequest):
    """
    Generate an example Excel file with user-provided reportName and reportId.
    Returns different example data each time.

    Args:
        request: Contains reportName and reportId

    Returns:
        Excel file as download
    """
    # Validate reportId is exactly 6 digits
    if not request.reportId.isdigit() or len(request.reportId) != 6:
        raise HTTPException(
            status_code=400,
            detail="Report ID must be exactly 6 digits (e.g., 200012, 500001)"
        )

    # Sample data pools for variety
    domains = ['finance', 'sales', 'marketing', 'operations', 'hr']
    categories = ['transactions', 'metrics', 'analytics', 'reports', 'dashboard']

    # Generate random formatSpecId and domainId (24-char hex)
    format_spec_id = ''.join(random.choices('0123456789abcdef', k=24))
    domain_id = ''.join(random.choices('0123456789abcdef', k=24))

    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report Metadata"

    # Header style
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    # Report metadata section
    ws['A1'] = 'reportName'
    ws['B1'] = request.reportName
    ws['A1'].fill = header_fill
    ws['A1'].font = header_font

    ws['A2'] = 'reportId'
    ws['B2'] = request.reportId
    ws['A2'].fill = header_fill
    ws['A2'].font = header_font

    ws['A3'] = 'reportType'
    ws['B3'] = random.choice(['analytics', 'operational', 'financial', 'marketing', 'sales'])
    ws['A3'].fill = header_fill
    ws['A3'].font = header_font

    ws['A4'] = 'formatSpecId'
    ws['B4'] = format_spec_id
    ws['A4'].fill = header_fill
    ws['A4'].font = header_font

    ws['A5'] = 'domainId'
    ws['B5'] = domain_id
    ws['A5'].fill = header_fill
    ws['A5'].font = header_font

    # Parameters section
    ws['A8'] = 'Parameters'
    ws['A8'].fill = header_fill
    ws['A8'].font = header_font

    ws['A9'] = 'parameterName'
    ws['B9'] = 'originalColumnName'
    ws['C9'] = 'parameterType'
    for col in ['A9', 'B9', 'C9']:
        ws[col].fill = header_fill
        ws[col].font = header_font

    # Generate 4-6 random parameters
    param_examples = [
        ('startDate', 'start_date', 'date'),
        ('endDate', 'end_date', 'date'),
        ('reportDate', 'report_date', 'timestamp'),
        ('userId', 'user_id', 'varchar(64)'),
        ('accountId', 'account_id', 'varchar(64)'),
        ('departmentId', 'department_id', 'varchar(64)'),
        ('regionCode', 'region_code', 'varchar(32)'),
        ('fiscalYear', 'fiscal_year', 'integer'),
        ('currencyCode', 'currency_code', 'char(3)'),
        ('statusFilter', 'status_filter', 'varchar(16)'),
    ]
    num_params = random.randint(4, 6)
    random.shuffle(param_examples)

    row_idx = 10
    for i in range(num_params):
        param_data = param_examples[i]
        ws[f'A{row_idx}'] = param_data[0]
        ws[f'B{row_idx}'] = param_data[1]
        ws[f'C{row_idx}'] = param_data[2]
        row_idx += 1

    # Columns section
    row_idx += 3
    ws[f'A{row_idx}'] = 'column name'
    ws[f'B{row_idx}'] = 'actual name'
    ws[f'C{row_idx}'] = 'type (postgresql)'
    ws[f'D{row_idx}'] = 'description'
    for col in [f'A{row_idx}', f'B{row_idx}', f'C{row_idx}', f'D{row_idx}']:
        ws[col].fill = header_fill
        ws[col].font = header_font

    # Generate exactly 25 columns
    column_examples = [
        ('id', 'ID', 'varchar(64)', 'unique identifier of the record'),
        ('name', 'Name', 'varchar(128)', 'display name of the entity'),
        ('created_at', 'Created At', 'timestamp', 'timestamp when the record was created'),
        ('updated_at', 'Updated At', 'timestamp', 'timestamp when the record was last updated'),
        ('status', 'Status', 'varchar(16)', 'current status of the record'),
        ('amount', 'Amount', 'numeric(18,2)', 'monetary amount in base currency'),
        ('quantity', 'Quantity', 'integer', 'number of items or units'),
        ('description', 'Description', 'varchar(256)', 'detailed description or notes'),
        ('category', 'Category', 'varchar(32)', 'classification category code'),
        ('user_id', 'User ID', 'varchar(64)', 'identifier of the associated user'),
        ('customer_id', 'Customer ID', 'varchar(64)', 'unique identifier of the customer'),
        ('email', 'Email', 'varchar(255)', 'email address of the user or customer'),
        ('phone', 'Phone', 'varchar(32)', 'phone number in international format'),
        ('address', 'Address', 'varchar(256)', 'full address or street address'),
        ('city', 'City', 'varchar(64)', 'city name'),
        ('country', 'Country', 'varchar(64)', 'country name or ISO code'),
        ('postal_code', 'Postal Code', 'varchar(16)', 'postal or ZIP code'),
        ('currency', 'Currency', 'char(3)', 'ISO 4217 currency code'),
        ('tax_amount', 'Tax Amount', 'numeric(18,2)', 'tax amount included or applied'),
        ('discount_amount', 'Discount Amount', 'numeric(18,2)', 'discount amount applied to the transaction'),
        ('total_amount', 'Total Amount', 'numeric(18,2)', 'final total amount after all adjustments'),
        ('payment_method', 'Payment Method', 'varchar(32)', 'payment method used'),
        ('transaction_date', 'Transaction Date', 'timestamp', 'date and time of the transaction'),
        ('reference_number', 'Reference Number', 'varchar(64)', 'reference or confirmation number'),
        ('notes', 'Notes', 'text', 'additional notes or comments'),
    ]

    num_columns = 25
    random.shuffle(column_examples)

    row_idx += 1
    for i in range(num_columns):
        col_data = column_examples[i]
        ws[f'A{row_idx}'] = col_data[0]
        ws[f'B{row_idx}'] = col_data[1]
        ws[f'C{row_idx}'] = col_data[2]
        ws[f'D{row_idx}'] = col_data[3]
        row_idx += 1

    # Adjust column widths
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 50

    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"example_{request.reportId}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
